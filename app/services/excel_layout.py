"""Private, bounded XLSX layout preview for the customer selecting a table.

This endpoint never sends spreadsheet cells to AI. Only a small redacted
preview is returned to the same browser, with no persistent file storage.
"""
from __future__ import annotations

import re
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from fastapi import UploadFile

from app.services.layout_intelligence import detect_workbook_tables
from app.services.user_import import MAX_XLSX_EXPANDED_BYTES, CustomerUploadError
from app.tools.preflight import MAX_COLUMNS, MAX_FILE_SIZE_BYTES

MAX_SHEETS = 30
MAX_PREVIEW_ROWS = 25
MAX_PREVIEW_VALUES_PER_ROW = 12

# Redact likely contact values in a row preview without hiding the headers.
SENSITIVE_PREVIEW = re.compile(
    r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b"
    r"|(?<!\d)(?:\+?1[.\s-]?)?\(?\d{3}\)?[.\s-]?\d{3}[.\s-]?\d{4}(?!\d)",
    re.IGNORECASE,
)


def _preview_text(value: object) -> str:
    if value is None:
        return ""
    safe = str(value).strip()
    if SENSITIVE_PREVIEW.search(safe):
        return "[sensitive value]"
    return safe[:67] + "..." if len(safe) > 70 else safe


def _validate_xlsx_archive(source: Path) -> None:
    try:
        with ZipFile(source) as archive:
            files = archive.infolist()
            if len(files) > 1500 or sum(f.file_size for f in files) > MAX_XLSX_EXPANDED_BYTES:
                raise CustomerUploadError("This Excel workbook exceeds safe processing limits.")
            if any(
                f.file_size > MAX_XLSX_EXPANDED_BYTES // 2
                or (f.compress_size == 0 and f.file_size > 0)
                or (f.compress_size > 0 and f.file_size > 150 * f.compress_size)
                for f in files
            ):
                raise CustomerUploadError("This Excel workbook has unsafe compression.")
    except (OSError, BadZipFile) as exc:
        raise CustomerUploadError("This Excel workbook is unreadable.") from exc


def inspect_xlsx_layout(source: Path) -> dict[str, object]:
    """Return sheet names, approximate dimensions, and 25 redacted row summaries."""
    _validate_xlsx_archive(source)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise CustomerUploadError("Excel support is unavailable.") from exc

    workbook = None
    try:
        workbook = load_workbook(
            source, read_only=True, data_only=True, keep_links=False
        )
        sheets = workbook.worksheets
        if not sheets or len(sheets) > MAX_SHEETS:
            raise CustomerUploadError("The workbook has no sheets or too many sheets.")

        result: list[dict[str, object]] = []
        for worksheet in sheets:
            preview_rows: list[dict[str, object]] = []
            for row_number, cells in enumerate(
                worksheet.iter_rows(max_row=MAX_PREVIEW_ROWS, max_col=MAX_COLUMNS), start=1
            ):
                samples: list[str] = []
                present = 0
                for cell in cells:
                    value = cell.value
                    if value is not None and str(value).strip():
                        present += 1
                        if len(samples) < MAX_PREVIEW_VALUES_PER_ROW:
                            samples.append(_preview_text(value))
                if present:
                    preview_rows.append({
                        "row_number": row_number,
                        "populated_cells": present,
                        "values": samples,
                    })
            result.append({
                "name": worksheet.title,
                "approximate_rows": worksheet.max_row,
                "approximate_columns": worksheet.max_column,
                "preview_rows": preview_rows,
            })
        inferred = detect_workbook_tables(sheets)
        return {
            "sheets": result,
            "default_sheet": (
                inferred["recommended_sheet"]
                if inferred["auto_selected"] else result[0]["name"]
            ),
            "default_header_row": (
                inferred["recommended_header_row"]
                if inferred["auto_selected"] else 1
            ),
            "intelligence": inferred,
            "no_model_call": True,
        }
    except CustomerUploadError:
        raise
    except Exception as exc:
        raise CustomerUploadError("Excel layout inspection failed safely.") from exc
    finally:
        if workbook is not None:
            workbook.close()


async def inspect_customer_workbook(
    upload: UploadFile, directory: Path
) -> dict[str, object]:
    """Bounded upload of an XLSX file exclusively for layout inspection."""
    if not upload.filename or not upload.filename.lower().endswith(".xlsx"):
        raise CustomerUploadError("Choose an XLSX workbook to select a worksheet.")
    path = directory / "layout.xlsx"
    size = 0
    try:
        with path.open("xb") as output:
            while True:
                chunk = await upload.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > MAX_FILE_SIZE_BYTES:
                    raise CustomerUploadError("The Excel file exceeds 10 MB.")
                output.write(chunk)
    finally:
        await upload.close()
    if size == 0:
        raise CustomerUploadError("The Excel file is empty.")
    return inspect_xlsx_layout(path)

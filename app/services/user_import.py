"""Bounded and private conversion of uploaded tabular data to CSV.

Only normalized CSV enters DataReady's existing governed engine. Files are
created inside a per-request TemporaryDirectory and never persisted or logged.
"""
from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from itertools import zip_longest
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from fastapi import UploadFile

from app.tools.preflight import (
    MAX_CELL_LENGTH,
    MAX_COLUMNS,
    MAX_FILE_SIZE_BYTES,
    MAX_ROWS,
)

SUPPORTED_EXTENSIONS = frozenset({".csv", ".json", ".xlsx"})
MAX_XLSX_EXPANDED_BYTES = 64 * 1024 * 1024


class CustomerUploadError(ValueError):
    """User-provided data is unsupported or exceeds the agreed safety limits."""


@dataclass
class ImportDiagnostics:
    """Disclosure of spreadsheet conversion choices for this one request."""

    warnings: list[str] = field(default_factory=list)
    uncached_formulas: int = 0
    cached_formulas: int = 0
    selected_sheet: str | None = None
    selected_header_row: int = 1


def _scalar(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        result = "true" if value else "false"
    elif isinstance(value, (str, int, float)):
        result = str(value)
    else:
        raise CustomerUploadError(
            "Only scalar text, number, boolean, or null values are supported."
        )
    if len(result) > MAX_CELL_LENGTH:
        raise CustomerUploadError("A value exceeds the allowed cell length.")
    return result


def _reject_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    obj: dict[str, Any] = {}
    for key, value in pairs:
        if key in obj:
            raise CustomerUploadError("JSON objects cannot contain duplicate keys.")
        obj[key] = value
    return obj


def _convert_json(source: Path, destination: Path) -> None:
    try:
        payload = json.loads(
            source.read_text(encoding="utf-8-sig"),
            object_pairs_hook=_reject_duplicate_keys,
            parse_constant=lambda value: (_ for _ in ()).throw(
                CustomerUploadError("Non-finite numeric JSON values are unsupported.")
            ),
        )
    except (ValueError, UnicodeError, json.JSONDecodeError) as error:
        raise CustomerUploadError("The uploaded JSON could not be safely parsed.") from error

    if isinstance(payload, dict) and set(payload) == {"records"}:
        payload = payload["records"]
    if not isinstance(payload, list) or not payload:
        raise CustomerUploadError(
            "JSON must be a nonempty array of flat objects or {\"records\": [...]}."
        )
    if len(payload) > MAX_ROWS:
        raise CustomerUploadError("The JSON dataset exceeds the row limit.")

    headers: list[str] = []
    for record in payload:
        if not isinstance(record, dict) or not record:
            raise CustomerUploadError("Every JSON row must be a nonempty flat object.")
        for key, value in record.items():
            if not isinstance(key, str) or not key or len(key) > MAX_CELL_LENGTH:
                raise CustomerUploadError("JSON column names must be nonempty text.")
            _scalar(value)
            if key not in headers:
                headers.append(key)
                if len(headers) > MAX_COLUMNS:
                    raise CustomerUploadError("The JSON dataset exceeds the column limit.")

    with destination.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for record in payload:
            writer.writerow([_scalar(record.get(key)) for key in headers])
    if destination.stat().st_size > MAX_FILE_SIZE_BYTES:
        raise CustomerUploadError("The converted CSV exceeds the safety size limit.")


def _convert_excel(
    source: Path,
    destination: Path,
    diagnostics: ImportDiagnostics | None = None,
    *,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> None:
    """Read saved formula results, never evaluate or execute Excel formulas.

    Where a formula has no cached result, preserve a visible inert placeholder
    and prohibit AI planning/export until the workbook is recalculated.
    """
    try:
        with ZipFile(source) as archive:
            infos = archive.infolist()
            expanded = sum(item.file_size for item in infos)
            if len(infos) > 1500 or expanded > MAX_XLSX_EXPANDED_BYTES:
                raise CustomerUploadError("The workbook exceeds safe processing limits.")
            if any(
                item.file_size > MAX_XLSX_EXPANDED_BYTES // 2
                or (item.compress_size == 0 and item.file_size > 0)
                or (item.compress_size > 0 and item.file_size > item.compress_size * 150)
                for item in infos
            ):
                raise CustomerUploadError("The workbook has unsafe compression.")
    except (BadZipFile, OSError) as error:
        raise CustomerUploadError("The XLSX workbook is unreadable.") from error

    if not 1 <= header_row <= MAX_ROWS + 1:
        raise CustomerUploadError("The selected Excel header row is not supported.")

    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise CustomerUploadError("Excel support is not installed on this server.") from error

    formulas = None
    calculated = None
    cached_count = 0
    unresolved_count = 0
    try:
        formulas = load_workbook(
            source, read_only=True, data_only=False, keep_links=False
        )
        calculated = load_workbook(
            source, read_only=True, data_only=True, keep_links=False
        )
        if not formulas.worksheets or not calculated.worksheets:
            raise CustomerUploadError("The workbook contains no worksheets.")
        if sheet_name is not None:
            if sheet_name not in formulas.sheetnames:
                raise CustomerUploadError("The selected Excel worksheet was not found.")
            worksheet = formulas[sheet_name]
            cached_sheet = calculated[sheet_name]
        else:
            worksheet = formulas.worksheets[0]
            cached_sheet = calculated.worksheets[0]
        if worksheet.max_column and worksheet.max_column > MAX_COLUMNS:
            raise CustomerUploadError("The worksheet exceeds the column limit.")
        if worksheet.max_row and worksheet.max_row > MAX_ROWS + header_row:
            raise CustomerUploadError("The worksheet exceeds the row limit.")

        if diagnostics is not None:
            diagnostics.selected_sheet = worksheet.title
            diagnostics.selected_header_row = header_row
            if sheet_name is not None or header_row != 1:
                diagnostics.warnings.append(
                    f"Excel table: worksheet {worksheet.title!r}, header at row "
                    f"{header_row}. Rows above the header were excluded."
                )

        with destination.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            for count, cell_pair in enumerate(
                zip_longest(worksheet.iter_rows(), cached_sheet.iter_rows()),
                start=1,
            ):
                formula_cells, cached_cells = cell_pair
                if count < header_row:
                    continue
                if count - header_row > MAX_ROWS:
                    raise CustomerUploadError("The worksheet exceeds the row limit.")
                if formula_cells is None or cached_cells is None:
                    raise CustomerUploadError("Workbook formula and value views differ.")
                if len(formula_cells) != len(cached_cells) or len(formula_cells) > MAX_COLUMNS:
                    raise CustomerUploadError("The worksheet exceeds safe column limits.")
                row: list[str] = []
                for cell, cached_cell in zip(formula_cells, cached_cells, strict=True):
                    if getattr(cell, "data_type", None) == "f":
                        if cached_cell.value is None:
                            unresolved_count += 1
                            value = "[UNEVALUATED_EXCEL_FORMULA]"
                        else:
                            cached_count += 1
                            value = cached_cell.value
                    else:
                        value = cell.value
                    if value is not None and not isinstance(value, (str, int, float, bool)):
                        value = str(value)
                    row.append(_scalar(value))
                writer.writerow(row)

        if not destination.is_file() or destination.stat().st_size == 0:
            raise CustomerUploadError("The selected Excel header row contains no table.")
        if destination.stat().st_size > MAX_FILE_SIZE_BYTES:
            raise CustomerUploadError("The converted worksheet exceeds the size limit.")

        if diagnostics is not None:
            diagnostics.cached_formulas += cached_count
            diagnostics.uncached_formulas += unresolved_count
            if cached_count:
                diagnostics.warnings.append(
                    f"{cached_count} Excel formula cell(s) were read from their saved results. "
                    "DataReady did not run or recalculate formulas; results may be stale."
                )
            if unresolved_count:
                diagnostics.warnings.append(
                    f"{unresolved_count} Excel formula cell(s) have no saved calculated value. "
                    "DataReady marked them as unevaluated. Audit can continue, but AI "
                    "planning and repaired-file export are disabled until the workbook "
                    "is recalculated and saved with results in Excel or another spreadsheet app."
                )
        elif unresolved_count:
            raise CustomerUploadError(
                "Excel formula results are missing. Recalculate and save the workbook."
            )
    except CustomerUploadError:
        raise
    except Exception as error:
        raise CustomerUploadError("The XLSX workbook could not be safely read.") from error
    finally:
        if formulas is not None:
            formulas.close()
        if calculated is not None:
            calculated.close()


async def receive_customer_dataset(
    upload: UploadFile,
    directory: Path,
    *,
    diagnostics: ImportDiagnostics | None = None,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> Path:
    """Receive a bounded upload, normalize if needed, and return a local CSV.

    Never use the supplied filename as a filesystem path. All work is isolated
    in a request-specific temporary directory. Caller owns its lifetime.
    """
    filename = upload.filename or ""
    if not filename or len(filename) > 255:
        raise CustomerUploadError("Select a CSV, JSON, or XLSX dataset.")
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise CustomerUploadError("Accepted formats: .csv, .json, and .xlsx.")

    incoming = directory / ("original" + extension)
    size = 0
    try:
        with incoming.open("xb") as handle:
            while True:
                data = await upload.read(1024 * 1024)
                if not data:
                    break
                size += len(data)
                if size > MAX_FILE_SIZE_BYTES:
                    raise CustomerUploadError("Upload exceeds the 10 MB file limit.")
                handle.write(data)
    finally:
        await upload.close()

    if size == 0:
        raise CustomerUploadError("The uploaded file is empty.")

    if extension != ".xlsx" and (sheet_name is not None or header_row != 1):
        raise CustomerUploadError(
            "Worksheet and header-row settings are available for Excel files only."
        )

    if extension == ".csv":
        return incoming

    destination = directory / "normalized.csv"
    if extension == ".json":
        _convert_json(incoming, destination)
    else:
        _convert_excel(
            incoming, destination, diagnostics,
            sheet_name=sheet_name, header_row=header_row,
        )
    return destination

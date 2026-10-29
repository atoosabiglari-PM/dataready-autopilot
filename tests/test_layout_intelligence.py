"""Automatic Excel table recognition: synthetic but realistically structured workbooks.

No private customer data and no NVIDIA inference are used in these tests.
"""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from app import live_product
from app.services.layout_intelligence import detect_workbook_tables

client = TestClient(live_product.app)


def hotel_directory_workbook():
    """A hotel-directory worksheet with seven introductory notes like the real example."""
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Hotel directory"
    sheet.append([
        "MANGHAL | HOTEL & MOTEL PARTNERSHIP DIRECTORY", None, None, None,
        "Center latitude", "Center longitude", "Earth radius (mi)",
        "Radius (mi)", "Boundary buffer (mi)",
    ] + [None] * 15)
    sheet.append([
        "Coverage: five miles", None, None, None,
        37.401817, -122.02755, 3958.7613, 5, .1,
    ] + [None] * 15)
    for sentence in [
        "Directory notes and scope",
        "Published locations are approximate",
        "Use filters to select your city",
        "Public-source research, not a complete inventory",
        "Scope: hotels, motels, extended-stay properties",
    ]:
        sheet.append([sentence] + [None] * 23)

    headers = [
        "ID", "Hotel / motel", "Street address", "City", "State", "ZIP",
        "Approx. miles", "Google Maps", "Public phone", "Public email",
        "Manager / contact", "Contact status",
    ] + [f"Field {i}" for i in range(12)]
    sheet.append(headers)
    for index in range(86):
        sheet.append([
            f"H{index + 1:03d}", f"Hotel {index + 1}",
            f"Address {index + 1}", "Sunnyvale", "CA", "94089",
            index / 10, "https://example.com", "NU", "NU", "NU",
            "Not contacted",
        ] + [f"value_{i}" for i in range(8)] + [None] * 4)
    return wb


def to_bytes(workbook):
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()


def test_numbered_record_values_are_not_scored_as_column_headers():
    from app.services.layout_intelligence import _label_score

    # These are record values, despite containing vocabulary such as
    # "hotel", "address", or "value" that can also appear in real headers.
    assert _label_score("Hotel 1") < _label_score("Hotel / motel")
    assert _label_score("Address 14") < _label_score("Street address")
    assert _label_score("value_0") < _label_score("Contact status")


def test_manghal_style_notes_are_not_misread_as_header():
    workbook = hotel_directory_workbook()
    detection = detect_workbook_tables(workbook.worksheets)
    assert detection["no_ai_call"] is True
    assert detection["confidence"] == "high"
    assert detection["auto_selected"] is True
    assert detection["recommended_sheet"] == "Hotel directory"
    assert detection["recommended_header_row"] == 8
    assert detection["leading_note_rows"] == 7
    assert detection["candidates"][0]["column_count"] == 24


def test_auto_detection_applies_to_actual_import_without_manual_worksheet_settings():
    raw = to_bytes(hotel_directory_workbook())
    response = client.post(
        "/api/audit", files={"file": ("hotels.xlsx", raw)}
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["model_invoked"] is False
    assert data["original_audit"]["row_count"] == 86
    assert data["original_audit"]["column_count"] == 24
    assert data["import_layout"] == {
        "sheet": "Hotel directory",
        "header_row": 8,
        "leading_note_rows_skipped": 7,
        "automatically_detected": True,
        "confidence": "high",
    }
    assert all(
        item["code"] != "SUSPICIOUS_COLUMN_NAME"
        for item in data["original_audit"]["issues"]
    )
    assert any("automatically recognized" in note.lower()
               for note in data["import_warnings"])


def test_multiple_plausible_worksheets_require_customer_choice():
    from openpyxl import Workbook

    wb = Workbook()
    first = wb.active
    first.title = "First"
    first.append(["product", "amount"])
    first.append(["A", 1])
    first.append(["B", 2])
    second = wb.create_sheet("Second")
    second.append(["sku", "quantity"])
    second.append(["A", 3])
    second.append(["B", 4])

    detection = detect_workbook_tables(wb.worksheets)
    assert detection["confidence"] == "review"
    assert detection["auto_selected"] is False
    assert detection["recommended_sheet"] is None
    assert len(detection["candidates"]) >= 2

    blob = to_bytes(wb)
    attempt = client.post("/api/audit", files={"file": ("two-tables.xlsx", blob)})
    assert attempt.status_code == 422
    assert "Select the worksheet and header row" in attempt.json()["detail"]

    selected = client.post(
        "/api/audit",
        files={"file": ("two-tables.xlsx", blob)},
        data={"sheet_name": "Second", "header_row": "1"},
    )
    assert selected.status_code == 200, selected.text
    layout = selected.json()["import_layout"]
    assert layout["sheet"] == "Second"
    assert layout["automatically_detected"] is False


def test_no_table_in_notes_only_workbook_is_not_auto_assumed():
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.append(["Instructions"])
    sheet.append(["This is a notes-only workbook."])
    sheet.append(["No table is present"])
    detection = detect_workbook_tables(wb.worksheets)
    assert detection["auto_selected"] is False
    assert detection["confidence"] == "unresolved"
    response = client.post(
        "/api/audit", files={"file": ("notes.xlsx", to_bytes(wb))}
    )
    assert response.status_code == 422


def test_preview_returns_recognition_without_exposing_customer_values():
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Client table"
    sheet.append(["contact", "email"])
    sheet.append(["Mr Example", "private@example.org"])
    sheet.append(["Ms Example", "private2@example.org"])
    raw = to_bytes(wb)
    response = client.post(
        "/api/workbook-layout", files={"file": ("contacts.xlsx", raw)}
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["intelligence"]["no_ai_call"] is True
    assert data["intelligence"]["method"] == "deterministic_structure_analysis"
    assert "private@example.org" not in response.text
    assert "private2@example.org" not in response.text
    assert "Mr Example" not in str(data["intelligence"])


def test_previously_correct_manual_layout_is_still_supported():
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Sheet1"
    sheet.append(["Explanation"])
    sheet.append(["id", "status"])
    sheet.append(["A", "ready"])
    sheet.append(["B", "ready"])
    response = client.post(
        "/api/audit",
        files={"file": ("inventory.xlsx", to_bytes(wb))},
        data={"sheet_name": "Sheet1", "header_row": "2"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["original_audit"]["row_count"] == 2
    assert response.json()["import_layout"]["automatically_detected"] is False


def test_frontend_includes_auto_detect_and_manual_override():
    response = client.get("/")
    assert response.status_code == 200
    assert "Auto-detected" in response.text
    assert "Multiple possible tables" in response.text
    assert 'id="excel-sheet"' in response.text
    assert 'id="excel-header-row"' in response.text
